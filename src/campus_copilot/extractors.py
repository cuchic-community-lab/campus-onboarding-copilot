import hashlib
import json
import re
from pathlib import Path
from typing import Dict, List, Tuple
from urllib.parse import quote, urljoin

from .config import FILES_DIR, FILES_JSON, QA_DIR, QRCODES_DIR, QRCODES_JSON, SITE_ROOT
from .models import DocumentRecord
from .source_policy import classify_file, classify_link, faq_decision


def stable_id(prefix: str, value: str) -> str:
    return prefix + "-" + hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalize_document_metadata(record: DocumentRecord) -> None:
    sample = (record.content or "")[:3000]
    if "学生手册" in record.title:
        record.published_at = None
        record.date_status = "mixed_compendium"
    spaced_date = re.compile(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")
    dates = spaced_date.findall(sample)
    if dates and record.authority_tier in {"official_policy", "official_guidance"} and record.date_status != "mixed_compendium":
        year, month, day = dates[-1]
        record.published_at = f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
        record.date_status = "document_date"
    elif record.authority_tier == "official_guidance" and record.date_status != "mixed_compendium":
        month_match = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月", (record.title + " " + sample[:800]))
        if month_match:
            record.published_at = f"{int(month_match.group(1)):04d}-{int(month_match.group(2)):02d}"
            record.date_status = "document_month"
    academic = re.search(r"(20\d{2})\s*[-—至]\s*(20\d{2})\s*学年", record.title + " " + sample[:1000])
    if academic:
        record.academic_year = f"{academic.group(1)}-{academic.group(2)}"
    title_and_sample = record.title + " " + sample[:1200]
    undergraduate_hits = len(re.findall(r"本科", title_and_sample))
    graduate_hits = len(re.findall(r"研究生", title_and_sample))
    if undergraduate_hits and undergraduate_hits >= graduate_hits:
        record.student_level = "undergraduate"
    elif graduate_hits:
        record.student_level = "graduate"
    elif "新生" in record.tags:
        record.student_level = "undergraduate"


def _extract_pdf(path: Path) -> Tuple[str, List[Dict[str, object]], str]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return "", [], "metadata_only_missing_pypdf"
    try:
        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            pages.append({"page_number": index, "text": text})
        content = "\n\n".join(str(page["text"]) for page in pages if page["text"])
        if not content:
            return "", pages, "ocr_required"
        return content, pages, "parsed"
    except Exception:
        return "", [], "parse_error"


def _extract_docx(path: Path) -> Tuple[str, str]:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        return "", "metadata_only_missing_python_docx"
    try:
        doc = Document(str(path))
        text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        return text, "parsed" if text else "empty"
    except Exception:
        return "", "parse_error"


def _extract_xlsx(path: Path, privacy_risk: str) -> Tuple[str, str]:
    if privacy_risk != "low":
        return "", "excluded_privacy_review"
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return "", "metadata_only_missing_openpyxl"
    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
        lines: List[str] = []
        for sheet in book.worksheets:
            lines.append("工作表：" + sheet.title)
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if values:
                    lines.append(" | ".join(values))
                if len(lines) >= 500:
                    break
        return "\n".join(lines), "parsed" if lines else "empty"
    except Exception:
        return "", "parse_error"


def parse_qa_markdown(markdown: str) -> List[Dict[str, str]]:
    """Parse Q&A markdown into question/answer units.

    Supports both layouts used by the corpus:
    - qa/ 目录：``## Q1: 问题`` + 正文回答（category 默认 General）
    - 站点根目录「新生问答整理.md」：``## 分类`` + ``### 问句`` + 正文回答
    """
    units: List[Dict[str, str]] = []
    question = ""
    answer_lines: List[str] = []
    category = "General"

    def flush() -> None:
        nonlocal answer_lines
        if question:
            units.append({"category": category, "question": question, "answer": "\n".join(answer_lines).strip()})
        answer_lines = []

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if line.startswith("### "):
            # 根目录格式：### 问句 是问题；若上一个 ## 行后面没有任何正文，
            # 说明它只是分类标题（把它的文本提升为 category，不当作问题）。
            if question and not answer_lines:
                category = question
            else:
                flush()
            question = re.sub(r"^###\s+", "", line).strip()
        elif line.startswith("## "):
            flush()
            question = re.sub(r"^##\s+", "", line).strip()
        elif question and line not in {"---", ""}:
            answer_lines.append(line)
    flush()
    return units


_EASTER_EGG_FILENAME_MARKERS = ("joke", "彩蛋", "段子", "冷笑话")

# 站点根目录下的高价值问答语料（v1.3）：与 qa/ 目录格式一致，一并入库。
ROOT_FAQ_FILENAME = "新生问答整理.md"


def _is_easter_egg_md(path: Path) -> bool:
    """True for joke/easter-egg markdown that must never enter the index.

    Detects by filename (joke/彩蛋/段子/冷笑话) or by an explicit hint in the
    document title/description (first lines). Kept deliberately broad so a
    qa/joke.md or qa/新生问答_段子版（彩蛋）.md can never be retrieved.
    """
    name = (path.stem or "").lower()
    if any(marker in name for marker in _EASTER_EGG_FILENAME_MARKERS):
        return True
    try:
        head = path.read_text(encoding="utf-8")[:600]
    except OSError:
        return False
    lowered = head.lower()
    return any(marker in lowered for marker in ("彩蛋", "冷笑话", "段子"))


def _site_root_faq_files() -> List[Path]:
    """Q&A markdown files at the site root that share the qa/ format.

    v1.3: only the explicitly curated 「新生问答整理.md」 is included. Easter-egg
    files (新生问答_段子版（彩蛋）.md …) are excluded by the shared
    ``_is_easter_egg_md`` filter, so joke content never enters the index.
    """
    candidate = SITE_ROOT / ROOT_FAQ_FILENAME
    if candidate.is_file() and not _is_easter_egg_md(candidate):
        return [candidate]
    return []


def faq_md_files(path: Path = QA_DIR) -> List[Path]:
    """All non-easter-egg markdown files under qa/ + the root FAQ md (sorted)."""
    files: List[Path] = []
    if path.is_dir():
        files.extend(p for p in path.glob("*.md") if not _is_easter_egg_md(p))
    elif path.is_file() and not _is_easter_egg_md(path):
        files.append(path)
    files.extend(_site_root_faq_files())
    return sorted(set(files))


def load_faq_units(path: Path = QA_DIR) -> Tuple[str, List[Dict[str, str]]]:
    """Load Q&A markdown from the site's qa/ directory + root 「新生问答整理.md」.

    Easter-egg / joke files (joke.md, 段子版…) are excluded so joke content
    never enters the knowledge index nor the FAQ answer pool.
    """
    markdown_parts: List[str] = []
    units: List[Dict[str, str]] = []
    for md_file in faq_md_files(path):
        text = md_file.read_text(encoding="utf-8")
        markdown_parts.append(text)
        units.extend(parse_qa_markdown(text))
    return "\n\n".join(markdown_parts), units


def _relative_url(local_path: str) -> str:
    """Turn a site-root-relative local path into a serveable /path URL."""
    normalized = local_path.replace("\\", "/")
    if normalized.startswith("/"):
        return normalized
    return "/" + normalized


def _build_file_document(item: Dict[str, object]) -> DocumentRecord:
    title = str(item.get("name", "Untitled"))
    media_type = str(item.get("type", "unknown"))
    tags = list(item.get("tags", []))
    decision = classify_file(title, media_type, tags)
    relative = str(item.get("path", "")).replace("\\", "/")
    local = SITE_ROOT / relative
    source_url = urljoin("https://hic.zihuanana.top/", quote(relative, safe="/"))
    record = DocumentRecord(
        document_id=stable_id("doc", relative),
        title=title,
        source_url=source_url,
        local_path=relative,
        media_type=media_type,
        source_kind=decision.source_kind,
        authority_tier=decision.authority_tier,
        assertion_policy=decision.assertion_policy,
        issuer=decision.issuer,
        description=str(item.get("desc", "")),
        tags=tags,
        uploaded_at=item.get("uploadTime"),
        privacy_risk=decision.privacy_risk,
    )
    if local.is_file():
        record.checksum = sha256_file(local)
        if media_type == "pdf":
            record.content, record.pages, record.parse_status = _extract_pdf(local)
        elif media_type == "word":
            record.content, record.parse_status = _extract_docx(local)
        elif media_type == "excel":
            record.content, record.parse_status = _extract_xlsx(local, record.privacy_risk)
        elif media_type == "image":
            # Images are evidence (thumbnail-able) but not assertable until OCR.
            record.parse_status = "ocr_required"
            record.content = f"{title}\n{record.description}".strip()
        else:
            record.parse_status = "metadata_only"
    else:
        record.parse_status = "not_downloaded"
    _normalize_document_metadata(record)
    return record


def _build_link_document(link: Dict[str, object]) -> DocumentRecord:
    title = str(link.get("title", "Untitled link"))
    url = str(link.get("url", ""))
    tags = list(link.get("tags", []))
    decision = classify_link(title, url, tags)
    return DocumentRecord(
        document_id=stable_id("link", url),
        title=title,
        source_url=url,
        local_path="",
        media_type="link",
        source_kind=decision.source_kind,
        authority_tier=decision.authority_tier,
        assertion_policy=decision.assertion_policy,
        issuer=decision.issuer,
        description=str(link.get("desc", "")),
        tags=tags,
        parse_status="resource_only",
        content=(title + "\n" + str(link.get("desc", ""))).strip(),
        student_level="undergraduate" if "新生" in tags or "选课" in tags else "all",
    )


def _build_qrcode_documents() -> List[DocumentRecord]:
    """QR codes become navigational resources whose images are thumbnailable
    evidence. Each entry is one document + a small descriptive chunk."""
    if not QRCODES_JSON.is_file():
        return []
    try:
        payload = json.loads(QRCODES_JSON.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    documents: List[DocumentRecord] = []
    for item in payload.get("items", []):
        if not item.get("visible", True):
            continue
        title = str(item.get("title", "二维码"))
        image = str(item.get("image", ""))
        category = str(item.get("category", ""))
        desc = str(item.get("desc", ""))
        local = SITE_ROOT / image.replace("\\", "/")
        record = DocumentRecord(
            document_id=stable_id("qr", image or title),
            title=title,
            source_url=_relative_url(image) if image else "",
            local_path=image,
            media_type="qr",
            source_kind="qr_resource",
            authority_tier="resource",
            assertion_policy="navigation_only",
            issuer="Student community",
            description=desc,
            tags=["二维码"] + ([category] if category else []),
            parse_status="resource_only",
            content=f"标题：{title}\n分类：{category}\n描述：{desc}".strip(),
            student_level="all",
        )
        if local.is_file():
            record.checksum = sha256_file(local)
        documents.append(record)
    return documents


def load_documents() -> Tuple[List[DocumentRecord], List[Dict[str, str]]]:
    """Load the full corpus from the site root: files + links + qrcodes + QA.

    Returns (documents, faq_units). Every path is relative to SITE_ROOT so the
    same build works on the local checkout and on the 宝塔 production server.
    """
    documents: List[DocumentRecord] = []

    # files.json manifest: files[] and links[]
    manifest_files: List[Dict[str, object]] = []
    manifest_links: List[Dict[str, object]] = []
    if FILES_JSON.is_file():
        try:
            manifest = json.loads(FILES_JSON.read_text(encoding="utf-8"))
            manifest_files = manifest.get("files", [])
            manifest_links = manifest.get("links", [])
        except (OSError, json.JSONDecodeError):
            manifest_files, manifest_links = [], []
    for item in manifest_files:
        documents.append(_build_file_document(item))
    for link in manifest_links:
        documents.append(_build_link_document(link))

    # qrcodes.json → QR entries with thumbnailable images
    documents.extend(_build_qrcode_documents())

    # qa/ directory → FAQ units (Q&A markdown)
    markdown, faq_units = load_faq_units()
    faq = faq_decision()
    if faq_units:
        # Point the FAQ document at a real servable markdown file so citation
        # file_path opens instead of 404ing on the aggregate "qa" directory
        # path. Prefer the root 「新生问答整理.md」 (richest content, served via
        # /corpus/ by server.py); otherwise fall back to the first qa/ file
        # (server.py maps /qa/ → QA_DIR).
        root_faq_files = _site_root_faq_files()
        qa_files = faq_md_files()
        faq_file_path = ""
        if root_faq_files:
            try:
                faq_file_path = f"corpus/{root_faq_files[0].name}"
            except (ValueError, OSError):
                faq_file_path = f"corpus/{ROOT_FAQ_FILENAME}"
        elif qa_files:
            try:
                faq_file_path = qa_files[0].resolve().relative_to(SITE_ROOT.resolve()).as_posix()
            except ValueError:
                faq_file_path = "qa/" + qa_files[0].name
        documents.append(DocumentRecord(
            document_id=stable_id("faq", "site-qa"),
            title="新生常见问题问答库",
            source_url="https://hic.zihuanana.top/#qa",
            local_path=faq_file_path or "qa",
            media_type="markdown",
            source_kind=faq.source_kind,
            authority_tier=faq.authority_tier,
            assertion_policy=faq.assertion_policy,
            issuer=faq.issuer,
            description="师哥师姐经验问答；非官方表述",
            tags=["新生", "Q&A", "学生经验"],
            cohort="2026",
            student_level="undergraduate",
            parse_status="parsed",
            content=markdown,
            checksum=hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
        ))
    return documents, faq_units
